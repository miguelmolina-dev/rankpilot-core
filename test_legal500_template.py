import os
import docx
from src.io.docx_manager import assemble_submission, extract_text_from_docx

def test_template():
    submission_data = {
        "Identity": {
            "firm_name": "Test Firm",
            "country": "USA",
            "practice_area": "M&A"
        },
        "InterviewContact": {
            "name": "Jane Doe",
            "job_title": "Partner",
            "email": "jane@testfirm.com"
        },
        "DepartmentInfo": {
            "team_name_internal": "Corporate"
        },
        "HeadOfTeam": {
            "name": "John Smith",
            "location": "NY"
        },
        "Metrics": {
            "partners_count_50_percent_plus": "5"
        },
        "Narratives": {
            "what_sets_us_apart": "We are fast.",
            "initiatives_and_innovation": "AI",
            "rankings_feedback": "Great."
        },
        "clients": [
            {"active_key_client": "Client A", "is_publishable": True, "is_new_client": "Yes"},
            {"active_key_client": "Client B", "is_publishable": True, "is_new_client": "No"},
            {"active_key_client": "Client C", "is_publishable": False, "is_new_client": "Yes"}
        ],
        "individual_nominations": {
            "leading_partners": [
                {"name": "Partner A", "location": "NY", "ranked_in_previous_edition": "Yes", "supporting_information": "Info A"},
                {"name": "Partner B", "location": "LA", "ranked_in_previous_edition": "No", "supporting_information": "Info B"},
                {"name": "Partner C", "location": "SF", "ranked_in_previous_edition": "Yes", "supporting_information": "Info C"}
            ],
            "next_generation_partners": [
                {"name": "NGP A", "location": "NY", "supporting_information": "Info NGP A"},
                {"name": "NGP B", "location": "LA", "supporting_information": "Info NGP B"}
            ],
            "leading_associates": [
                {"name": "LA A", "location": "NY", "supporting_information": "Info LA A"},
                {"name": "LA B", "location": "LA", "supporting_information": "Info LA B"}
            ]
        },
        "ArrivalDeparture": [
            {"name": "Arr A", "position_role": "Partner", "action": "Joined"},
            {"name": "Dep B", "position_role": "Counsel", "action": "Departed"}
        ],
        "WorkHighlight": [
            {"publishable_summary": "Summary 1"},
            {"publishable_summary": "Summary 2"}
        ],
        "matters": [
            {
                "client_name": "Matter Client 1", "industry_sector": "Tech", "matter_description": "Desc 1", "is_publishable": True,
                "deal_value": "$100M", "jurisdictions_involved": "USA, UK",
                "leading_partners": [{"name": "M_P1", "office": "NY", "practice_area": "M&A"}, {"name": "M_P2", "office": "NY", "practice_area": "M&A"}],
                "other_partners": [{"name": "O_P1", "office": "NY", "practice_area": "M&A"}, {"name": "O_P2", "office": "NY", "practice_area": "M&A"}],
                "leading_associates": [{"name": "A_A1", "office": "NY", "practice_area": "M&A"}, {"name": "A_A2", "office": "NY", "practice_area": "M&A"}],
                "other_associates": [{"name": "O_A1", "office": "NY", "practice_area": "M&A"}, {"name": "O_A2", "office": "NY", "practice_area": "M&A"}],
                "external_firms_advising": [{"firm_name": "Rival Firm", "role_details": "Opposing", "entity_advised": "Target"}],
                "dates": {"start": "2024-01-01", "end": "2024-12-31"}
            },
            {
                "client_name": "Matter Client 2", "industry_sector": "Finance", "matter_description": "Desc 2", "is_publishable": True,
                "deal_value": "$200M", "jurisdictions_involved": "USA",
                "leading_partners": [{"name": "M_P3", "office": "NY", "practice_area": "M&A"}, {"name": "M_P4", "office": "NY", "practice_area": "M&A"}],
                "other_partners": [{"name": "O_P3", "office": "NY", "practice_area": "M&A"}, {"name": "O_P4", "office": "NY", "practice_area": "M&A"}],
                "leading_associates": [{"name": "A_A3", "office": "NY", "practice_area": "M&A"}, {"name": "A_A4", "office": "NY", "practice_area": "M&A"}],
                "other_associates": [{"name": "O_A3", "office": "NY", "practice_area": "M&A"}, {"name": "O_A4", "office": "NY", "practice_area": "M&A"}],
                "external_firms_advising": [{"firm_name": "Partner Firm", "role_details": "Co-counsel", "entity_advised": "Target"}],
                "dates": {"start": "2024-02-01", "end": "2024-11-30"}
            },
            {
                "client_name": "Matter Client 3", "industry_sector": "Health", "matter_description": "Desc 3", "is_publishable": False,
                "deal_value": "$300M", "jurisdictions_involved": "USA, EU",
                "leading_partners": [{"name": "M_P5", "office": "NY", "practice_area": "M&A"}, {"name": "M_P6", "office": "NY", "practice_area": "M&A"}],
                "other_partners": [{"name": "O_P5", "office": "NY", "practice_area": "M&A"}, {"name": "O_P6", "office": "NY", "practice_area": "M&A"}],
                "leading_associates": [{"name": "A_A5", "office": "NY", "practice_area": "M&A"}, {"name": "A_A6", "office": "NY", "practice_area": "M&A"}],
                "other_associates": [{"name": "O_A5", "office": "NY", "practice_area": "M&A"}, {"name": "O_A6", "office": "NY", "practice_area": "M&A"}],
                "external_firms_advising": [{"firm_name": "Other Firm", "role_details": "Local counsel", "entity_advised": "Buyer"}],
                "dates": {"start": "2024-03-01", "end": "2024-10-31"}
            }
        ]
    }

    template_path = 'templates/legal500-us-submissions-template-2026-1-1-2.docx'
    output_dir = 'test_output'
    os.makedirs(output_dir, exist_ok=True)

    print(f"Generating document with assemble_submission...")
    out_path = assemble_submission(template_path, output_dir, submission_data)
    print(f"Document successfully generated at {out_path}")

    # Inspect the generated document to verify loop duplication and tag replacements
    doc = docx.Document(out_path)

    publishable_matters = 0
    non_publishable_matters = 0

    for table in doc.tables:
        if len(table.rows) > 1 and len(table.columns) > 1:
            try:
                first_cell_text = table.cell(0, 0).text
            except Exception:
                continue

            if 'Publishable matter' in first_cell_text:
                publishable_matters += 1
                client = table.cell(2, 0).text.strip()
                print(f"Found Publishable Matter -> Client: {client}")
            elif 'Non-publishable matter' in first_cell_text:
                non_publishable_matters += 1
                client = table.cell(2, 0).text.strip()
                print(f"Found Non-publishable Matter -> Client: {client}")

    assert publishable_matters == 2, f"Expected 2 publishable matter tables, found {publishable_matters}"
    assert non_publishable_matters == 1, f"Expected 1 non-publishable matter table, found {non_publishable_matters}"

    print("All matter tables correctly duplicated and rendered!")

    print("\n--- Final Document Content ---")
    extracted_text = extract_text_from_docx(out_path)
    print(extracted_text)
    print("------------------------------")

if __name__ == '__main__':
    test_template()
